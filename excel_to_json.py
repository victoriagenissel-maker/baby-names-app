import json
from openpyxl import load_workbook

# Ouvre le fichier Excel
workbook = load_workbook("names.xlsx", read_only=True)
sheet = workbook.active

names = []

# On commence à la 2e ligne pour ignorer les titres
for row in sheet.iter_rows(min_row=2, values_only=True):

    if len(row) < 3:
        continue

    name = row[0]
    gender = row[1]
    origin = row[2]

    # Ignore les lignes sans prénom
    if not name:
        continue

    names.append({
        "name": str(name).strip(),
        "gender": str(gender).strip() if gender else "",
        "origin": str(origin).strip() if origin else ""
    })

# Création du fichier JSON
with open("names.json", "w", encoding="utf-8") as file:
    json.dump(names, file, ensure_ascii=False, indent=2)

print(f"{len(names)} prénoms ont été convertis.")